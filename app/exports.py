from __future__ import annotations

from io import BytesIO
from textwrap import wrap

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def performance_history_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Performance History"
    headers = ["Supplier", "Contract", "SO", "Call-up", "Project", "Type", "Region", "Date", "Evaluator", "Model", "Score", "Outcome", "Status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F62FE")
    for row in rows:
        ws.append([safe_cell(row.get(key, "")) for key in ("supplier", "contract", "standing_offer", "call_up", "project", "procurement_type", "region", "date", "evaluator", "model", "score", "outcome", "status")])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(42, max(12, max(len(str(c.value or "")) for c in column) + 2))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def tabular_report_pdf(title: str, subtitle: str, headers: list[str], rows: list[list]) -> bytes:
    out = BytesIO()
    pdf = canvas.Canvas(out, pagesize=LETTER)
    width, height = LETTER
    y = height - 48
    pdf.setFont("Helvetica-Bold", 15)
    pdf.drawString(42, y, title)
    y -= 18
    pdf.setFont("Helvetica", 8)
    pdf.drawString(42, y, subtitle)
    y -= 24
    widths = [max(70, int((width - 84) / max(1, len(headers)))) for _ in headers]
    pdf.setFont("Helvetica-Bold", 7)
    x = 42
    for i, header in enumerate(headers):
        pdf.drawString(x, y, str(header)[:18])
        x += widths[i]
    y -= 12
    pdf.line(42, y, width - 42, y)
    y -= 12
    pdf.setFont("Helvetica", 7)
    for row in rows:
        if y < 55:
            pdf.showPage(); y = height - 48; pdf.setFont("Helvetica", 7)
        x = 42
        for i, value in enumerate(row):
            pdf.drawString(x, y, str(value if value is not None else "")[:22])
            x += widths[i]
        y -= 12
    pdf.setTitle(title)
    pdf.save()
    return out.getvalue()


def correspondence_pdf(context: dict) -> bytes:
    out = BytesIO()
    pdf = canvas.Canvas(out, pagesize=LETTER)
    width, height = LETTER
    y = height - 60
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(54, y, context.get("organization") or "Contracting Organization")
    y -= 28
    pdf.setFont("Helvetica", 10)
    pdf.drawString(54, y, context["date"])
    y -= 28
    pdf.drawString(54, y, context["supplier"])
    y -= 16
    pdf.drawString(54, y, f"Contract: {context['contract']}")
    y -= 30
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(54, y, context["subject"])
    y -= 28
    pdf.setFont("Helvetica", 10)
    for paragraph in context["paragraphs"]:
        for line in wrap(paragraph, 92):
            pdf.drawString(54, y, line)
            y -= 14
            if y < 70:
                pdf.showPage(); y = height - 60; pdf.setFont("Helvetica", 10)
        y -= 10
    pdf.drawString(54, y, "Contracting Authority")
    pdf.setTitle(context["subject"])
    pdf.save()
    return out.getvalue()
