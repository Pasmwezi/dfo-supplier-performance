from pathlib import Path
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain import CONSTRUCTION_CRITERIA
from app.exports import performance_history_xlsx
from app.main import create_app


def make_client(tmp_path: Path):
    tokens = {
        "test-token": {"name": "Test Admin", "roles": ["ADMIN", "EVALUATOR", "REVIEWER", "APPROVER", "DECISION_MAKER"]},
        "review-token": {"name": "Independent Reviewer", "roles": ["REVIEWER"]},
        "approve-token": {"name": "Independent Approver", "roles": ["APPROVER"]},
        "decision-token": {"name": "Delegated Decision Maker", "roles": ["DECISION_MAKER"]},
    }
    app = create_app(f"sqlite:///{tmp_path / 'test.db'}", seed=False, auth_tokens=tokens)
    return TestClient(app, headers={"Authorization": "Bearer test-token"})


def test_supplier_profile_contract_counts_update_from_contract_records(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Multi-Contract Supplier"}).json()
    initial = client.get(f"/api/suppliers/{supplier['id']}/profile").json()
    assert initial["contract_summary"] == {"total": 0, "active": 0, "historical": 0}

    for number, status in (("MULTI-ACTIVE-1", "ACTIVE"), ("MULTI-CLOSED-2", "CLOSED")):
        response = client.post("/api/contracts", json={
            "supplier_id": supplier["id"], "contract_number": number,
            "procurement_type": "Construction", "region": "Atlantic", "status": status,
        })
        assert response.status_code == 201, response.text

    profile = client.get(f"/api/suppliers/{supplier['id']}/profile").json()
    assert profile["contract_summary"] == {"total": 2, "active": 1, "historical": 1}
    page = client.get("/suppliers").text
    assert f'data-supplier-id="{supplier["id"]}"' in page
    assert 'data-contract-total="2"' in page
    assert 'data-contract-active="1"' in page
    assert 'data-contract-historical="1"' in page


def test_health_and_dashboard(tmp_path):
    client = make_client(tmp_path)
    assert client.get("/health").json() == {"status": "ok"}
    response = client.get("/api/dashboard")
    assert response.status_code == 200
    assert set(response.json()) >= {"active_suppliers", "active_contracts", "average_performance", "risk_distribution"}


def test_default_app_does_not_seed_demo_records(tmp_path):
    app = create_app(
        f"sqlite:///{tmp_path / 'no-demo.db'}",
        auth_tokens={"test-token": {"name": "Test Admin", "roles": ["ADMIN"]}},
    )
    client = TestClient(app, headers={"Authorization": "Bearer test-token"})
    assert client.get("/api/suppliers").json() == []
    assert "/evaluations/new" not in client.get("/").text
    assert client.get("/evaluations/new").status_code == 403


def test_evaluation_captures_official_project_and_contract_information(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Project Information Contractor"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "PROJECT-DETAILS-1",
        "project_number": "PROJECT-26-001", "procurement_type": "Construction",
        "region": "Atlantic", "status": "ACTIVE",
    }).json()
    details = {
        "client_reference_number": "DFO-CLIENT-7788",
        "description_of_work": "Rehabilitation of the East Harbour breakwater.",
        "firm_address": "100 Harbour Road, Halifax, NS",
        "contractor_superintendent": "Jordan Chen",
        "project_manager_name": "Alex Martin",
        "project_manager_telephone": "902-555-0100",
        "project_manager_fax": "902-555-0101",
        "project_manager_cell": "902-555-0102",
        "project_manager_email": "alex.martin@example.gc.ca",
        "contract_award_amount": 1250000,
        "contract_award_date": "2026-01-15",
        "final_amount": 1275000,
        "contract_completion_date": "2026-06-30",
        "contract_changes_count": 3,
        "final_certificate_date": "2026-07-10",
    }
    response = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Project Evaluator",
        "scores": {name: 15 for name in CONSTRUCTION_CRITERIA}, "project_details": details,
    })
    assert response.status_code == 201, response.text
    evaluation = response.json()
    assert evaluation["project_details"] == details
    page = client.get(f"/evaluations/{evaluation['id']}/view").text
    assert "DFO-CLIENT-7788" in page
    assert "Rehabilitation of the East Harbour breakwater." in page
    assert "Jordan Chen" in page
    assert "$1,250,000.00" in page
    history = client.get(f"/api/evaluations/{evaluation['id']}/history").json()
    assert history["versions"][0]["snapshot"]["project_details"] == details

    revised_details = details | {"final_amount": 1280000, "contract_changes_count": 4}
    revised = client.patch(f"/api/evaluations/{evaluation['id']}", json={"project_details": revised_details})
    assert revised.status_code == 200, revised.text
    assert revised.json()["version"] == 2
    assert revised.json()["project_details"]["final_amount"] == 1280000
    revised_history = client.get(f"/api/evaluations/{evaluation['id']}/history").json()
    assert revised_history["versions"][-1]["snapshot"]["project_details"] == revised_details
    assert any(item["field_name"] == "project_details" for item in revised_history["audit"])


def test_new_evaluation_page_includes_official_project_information_fields(tmp_path):
    page = make_client(tmp_path).get("/evaluations/new").text
    for field in (
        "Client reference number", "Description of work", "Firm / contractor address",
        "Contractor’s superintendent", "Project manager name", "Telephone number",
        "Fax number", "Cell number", "Email address", "Contract award amount",
        "Contract award date", "Final amount", "Contract completion date",
        "Number of change orders / amendments", "Final certificate date",
    ):
        assert field in page


def test_create_supplier_contract_and_construction_evaluation(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Atlantic Builders Ltd.", "business_number": "BN-123"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "F5211-260001",
        "project_number": "DFO-001", "procurement_type": "Construction",
        "region": "Atlantic", "department": "Fisheries and Oceans Canada",
        "status": "ACTIVE"
    }).json()
    evaluation = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Test Evaluator",
        "scores": {
            "Quality of Workmanship": 17, "Time": 16, "Project Management": 15,
            "Contract Management": 14, "Health and Safety": 18
        },
        "comments": "Evidence reviewed.", "issues": ["schedule"]
    })
    assert evaluation.status_code == 201
    data = evaluation.json()
    assert data["total_score"] == 80
    assert data["outcome"] == "MEETS_EXPECTATIONS"
    assert data["status"] == "DRAFT"


def test_evaluator_can_create_supplier_contract_and_evaluation_together(tmp_path):
    client = make_client(tmp_path)
    response = client.post("/api/evaluations", json={
        "new_supplier": {"name": "New Harbour Contractors Ltd.", "business_number": "BN-NEW-001"},
        "new_contract": {
            "contract_number": "F5211-260099", "project_number": "ATL-NEW-01",
            "procurement_type": "Construction", "region": "Atlantic",
            "performance_regime": "GI16_GC1_22", "status": "ACTIVE"
        },
        "model": "CONSTRUCTION", "evaluator": "Test Evaluator",
        "scores": {
            "Quality of Workmanship": 17, "Time": 16, "Project Management": 15,
            "Contract Management": 14, "Health and Safety": 18
        },
        "comments": "First evaluation for a newly entered supplier."
    })
    assert response.status_code == 201
    evaluation = response.json()
    suppliers = client.get("/api/suppliers").json()
    contracts = client.get("/api/contracts").json()
    assert suppliers == [{"id": suppliers[0]["id"], "name": "New Harbour Contractors Ltd.", "business_number": "BN-NEW-001", "status": "ACTIVE"}]
    assert contracts[0]["supplier_id"] == suppliers[0]["id"]
    assert contracts[0]["contract_number"] == "F5211-260099"
    assert evaluation["contract_id"] == contracts[0]["id"]
    history = client.get(f"/api/evaluations/{evaluation['id']}/history").json()
    assert {entry["action"] for entry in history["audit"]} >= {"CREATE_SUPPLIER_WITH_EVALUATION", "CREATE_CONTRACT_WITH_EVALUATION", "CREATE"}


def test_new_evaluation_page_offers_inline_supplier_and_contract_entry(tmp_path):
    client = make_client(tmp_path)
    page = client.get("/evaluations/new")
    assert page.status_code == 200
    assert "Add a new firm / contractor / supplier" in page.text
    assert 'name="new_supplier_name"' in page.text
    assert 'name="new_contract_number"' in page.text
    assert 'name="new_procurement_type"' in page.text
    assert 'data-na-for="Project Management"' in page.text
    assert 'data-na-for="Contract Management"' in page.text
    assert 'data-na-for="Design"' in page.text
    assert 'data-na-for="Management"' in page.text
    assert 'data-na-for="Cost"' in page.text
    assert 'data-na-for="Quality of Workmanship"' not in page.text
    assert 'data-na-for="Quality of Results"' not in page.text


def test_cperf_optional_criteria_support_not_applicable_and_reject_required_na(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Applicability Test Firm"}).json()
    construction_contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "NA-CONSTRUCTION-001", "procurement_type": "Construction",
        "region": "Atlantic", "performance_regime": "GI16_GC1_22"
    }).json()
    ae_contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "NA-AE-001", "procurement_type": "Engineering Services",
        "region": "Atlantic", "performance_regime": "GI23_GC26_2913_1"
    }).json()
    construction_scores = {
        "Quality of Workmanship": 15, "Time": 15, "Project Management": None,
        "Contract Management": None, "Health and Safety": 15,
    }
    response = client.post("/api/evaluations", json={
        "contract_id": construction_contract["id"], "model": "CONSTRUCTION", "evaluation_date": "2026-07-19",
        "evaluator": "Evaluator A", "scores": construction_scores, "issues": [], "comments": "N/A reflects project applicability."
    })
    assert response.status_code == 201, response.text
    assert response.json()["total_score"] == 75.0
    assert response.json()["ratings"]["Project Management"] == "Not Applicable"
    detail = client.get(f"/evaluations/{response.json()['id']}/view")
    assert "Excluded (20%)" in detail.text
    assert "Not Applicable" in detail.text
    assert "45 / 60 applicable points" in detail.text

    no_rationale = client.post("/api/evaluations", json={
        "contract_id": construction_contract["id"], "model": "CONSTRUCTION", "evaluation_date": "2026-07-19",
        "evaluator": "Evaluator A", "scores": construction_scores, "issues": [], "comments": ""
    })
    assert no_rationale.status_code == 422
    assert "explain each Not Applicable" in no_rationale.json()["detail"]

    required_na = dict(construction_scores, **{"Quality of Workmanship": None, "Project Management": 15})
    response = client.post("/api/evaluations", json={
        "contract_id": construction_contract["id"], "model": "CONSTRUCTION", "evaluation_date": "2026-07-19",
        "evaluator": "Evaluator A", "scores": required_na, "issues": [], "comments": "Quality criterion was outside scope."
    })
    assert response.status_code == 422
    assert "cannot be Not Applicable" in response.json()["detail"]

    ae_scores = {"Design": None, "Quality of Results": 16, "Management": None, "Time": 14, "Cost": None}
    response = client.post("/api/evaluations", json={
        "contract_id": ae_contract["id"], "model": "AE_CPERF", "evaluation_date": "2026-07-19",
        "evaluator": "Evaluator A", "scores": ae_scores, "issues": [], "comments": "N/A criteria were outside project scope."
    })
    assert response.status_code == 201, response.text
    assert response.json()["total_score"] == 75.0
    assert response.json()["ratings"]["Cost"] == "Not Applicable"


def test_ae_rejects_invalid_weight_total(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Northstar Engineering", "business_number": "BN-456"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "F5211-260002",
        "procurement_type": "Engineering Services", "region": "Pacific",
        "department": "Fisheries and Oceans Canada", "status": "ACTIVE"
    }).json()
    response = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "AE", "evaluator": "Test Evaluator",
        "scores": {"Quality of Design": 15}, "weights": {"Quality of Design": 99}
    })
    assert response.status_code == 422


def test_cperf_rejects_modified_prescribed_weights(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "CPERF Supplier"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "CPERF-1",
        "procurement_type": "Engineering Services", "region": "Pacific",
        "performance_regime": "GI23_GC26_2913_1", "status": "ACTIVE"
    }).json()
    response = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "AE_CPERF", "evaluator": "Test Evaluator",
        "scores": {"Design": 20, "Quality of Results": 1, "Management": 1, "Time": 1, "Cost": 1},
        "weights": {"Design": 96, "Quality of Results": 1, "Management": 1, "Time": 1, "Cost": 1}
    })
    assert response.status_code == 422
    assert "prescribed" in response.json()["detail"].lower()


def test_update_creates_version_and_field_level_audit(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Audit Supplier", "business_number": "BN-789"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "F5211-260003",
        "procurement_type": "Construction", "region": "Central and Arctic",
        "department": "Fisheries and Oceans Canada", "status": "ACTIVE"
    }).json()
    evaluation = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator A",
        "scores": {"Quality of Workmanship": 12, "Time": 12, "Project Management": 12, "Contract Management": 12, "Health and Safety": 12}
    }).json()
    updated = client.patch(f"/api/evaluations/{evaluation['id']}", headers={"X-User": "Reviewer B"}, json={"comments": "Revised after evidence review."})
    assert updated.status_code == 200
    history = client.get(f"/api/evaluations/{evaluation['id']}/history").json()
    assert len(history["versions"]) == 2
    assert any(a["field_name"] == "comments" and a["user"].startswith("Test Admin [") for a in history["audit"])
    unchanged = client.patch(f"/api/evaluations/{evaluation['id']}", json={"comments": "Revised after evidence review."})
    assert unchanged.status_code == 200
    assert unchanged.json()["version"] == 2
    assert len(client.get(f"/api/evaluations/{evaluation['id']}/history").json()["versions"]) == 2


def test_update_rejects_null_for_required_evaluation_fields(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Null Guard Supplier"}).json()
    contract = client.post("/api/contracts", json={"supplier_id": supplier["id"], "contract_number": "NULL-1", "procurement_type": "Construction", "region": "Atlantic"}).json()
    evaluation = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator",
        "scores": {"Quality of Workmanship": 12, "Time": 12, "Project Management": 12, "Contract Management": 12, "Health and Safety": 12}
    }).json()
    for field in ("evaluation_date", "evaluator", "scores", "comments", "issues"):
        response = client.patch(f"/api/evaluations/{evaluation['id']}", json={field: None})
        assert response.status_code == 422, field


def test_approval_workflow_and_exports(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Export Supplier", "business_number": "BN-999"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "F5211-260004",
        "procurement_type": "Construction", "region": "Quebec",
        "department": "Fisheries and Oceans Canada", "status": "ACTIVE"
    }).json()
    evaluation = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator",
        "scores": {"Quality of Workmanship": 18, "Time": 18, "Project Management": 18, "Contract Management": 18, "Health and Safety": 18}
    }).json()
    action_tokens = {"submit": "test-token", "review": "review-token", "approve": "approve-token"}
    for action in ("submit", "review", "approve"):
        response = client.post(f"/api/evaluations/{evaluation['id']}/workflow/{action}", headers={"Authorization": f"Bearer {action_tokens[action]}"})
        assert response.status_code == 200
    assert response.json()["status"] == "APPROVED"
    assert client.get(f"/api/evaluations/{evaluation['id']}/correspondence.pdf").content.startswith(b"%PDF")
    assert "spreadsheetml" in client.get("/api/reports/performance-history.xlsx").headers["content-type"]
    assert client.get("/api/reports/supplier-risk.pdf").content.startswith(b"%PDF")


def test_suspension_decision_is_separate_from_system_recommendation(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Decision Supplier", "business_number": "BN-DEC"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "F5211-260005",
        "procurement_type": "Construction", "region": "National Capital",
        "department": "Fisheries and Oceans Canada", "status": "ACTIVE"
    }).json()
    evaluation = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator",
        "scores": {"Quality of Workmanship": 5, "Time": 5, "Project Management": 5, "Contract Management": 5, "Health and Safety": 5}
    }).json()
    profile = client.get(f"/api/suppliers/{supplier['id']}/profile").json()
    assert profile["active_suspension_decisions"] == []
    decision_payload = {
        "supplier_id": supplier["id"], "evaluation_id": evaluation["id"],
        "decision_type": "SUSPENSION", "status": "PENDING",
        "effective_date": "2026-07-17", "expiry_date": "2027-07-16",
        "rationale": "Decision made after notice, representations and delegated approval.",
        "authority": "Director General", "notice_date": "2026-06-01",
        "representation_deadline": "2026-06-20", "representations_summary": "Supplier representations reviewed and addressed.",
        "legal_review_reference": "LEGAL-2026-001", "delegated_authority_reference": "DOA-2026-001"
    }
    assert client.post("/api/decisions", json=decision_payload).status_code == 409
    for action, token in (("submit", "test-token"), ("review", "review-token"), ("approve", "approve-token")):
        assert client.post(f"/api/evaluations/{evaluation['id']}/workflow/{action}", headers={"Authorization": f"Bearer {token}"}).status_code == 200
    invalid_chronology = decision_payload | {"notice_date": "2026-07-18", "representation_deadline": "2026-07-19"}
    assert client.post("/api/decisions", json=invalid_chronology).status_code == 422
    decision = client.post("/api/decisions", json=decision_payload)
    assert decision.status_code == 201
    assert client.post(f"/api/decisions/{decision.json()['id']}/approve").status_code == 409
    decision = client.post(f"/api/decisions/{decision.json()['id']}/approve", headers={"Authorization": "Bearer decision-token"})
    assert decision.status_code == 200
    profile = client.get(f"/api/suppliers/{supplier['id']}/profile").json()
    assert len(profile["active_suspension_decisions"]) == 1


def test_write_requires_authentication(tmp_path):
    tokens = {"test-token": {"name": "Test Admin", "roles": ["ADMIN"]}}
    app = create_app(f"sqlite:///{tmp_path / 'auth.db'}", seed=False, auth_tokens=tokens)
    anonymous = TestClient(app)
    assert anonymous.post("/api/suppliers", json={"name": "Anonymous Supplier"}).status_code == 401
    assert anonymous.get("/api/suppliers").status_code == 401


def test_draft_correspondence_is_blocked(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Draft Supplier"}).json()
    contract = client.post("/api/contracts", json={"supplier_id": supplier["id"], "contract_number": "DRAFT-1", "procurement_type": "Construction", "region": "Atlantic", "status": "ACTIVE"}).json()
    evaluation = client.post("/api/evaluations", json={"contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator", "scores": {"Quality of Workmanship": 15, "Time": 15, "Project Management": 15, "Contract Management": 15, "Health and Safety": 15}}).json()
    assert client.get(f"/api/evaluations/{evaluation['id']}/correspondence.pdf").status_code == 409


def test_evaluation_detail_exposes_only_authorized_available_actions(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Role Aware Supplier"}).json()
    contract = client.post("/api/contracts", json={"supplier_id": supplier["id"], "contract_number": "ROLE-1", "procurement_type": "Construction", "region": "Atlantic", "status": "ACTIVE"}).json()
    evaluation = client.post("/api/evaluations", json={"contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator", "scores": {"Quality of Workmanship": 15, "Time": 15, "Project Management": 15, "Contract Management": 15, "Health and Safety": 15}}).json()
    page_url = f"/evaluations/{evaluation['id']}/view"

    evaluator_page = client.get(page_url).text
    assert 'data-action="submit"' in evaluator_page
    assert 'id="workflow-message"' in evaluator_page
    assert "Available after approval" in evaluator_page
    assert f'/api/evaluations/{evaluation["id"]}/correspondence.pdf' not in evaluator_page

    reviewer_page = client.get(page_url, headers={"Authorization": "Bearer review-token"}).text
    assert 'data-action="submit"' not in reviewer_page


def test_return_action_moves_from_reviewer_to_approver_after_review(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Return Role Supplier"}).json()
    contract = client.post("/api/contracts", json={"supplier_id": supplier["id"], "contract_number": "RETURN-1", "procurement_type": "Construction", "region": "Atlantic"}).json()
    evaluation = client.post("/api/evaluations", json={"contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator", "scores": {"Quality of Workmanship": 15, "Time": 15, "Project Management": 15, "Contract Management": 15, "Health and Safety": 15}}).json()
    page_url = f"/evaluations/{evaluation['id']}/view"
    assert client.post(f"/api/evaluations/{evaluation['id']}/workflow/submit").status_code == 200
    reviewer_headers = {"Authorization": "Bearer review-token"}
    reviewer_page = client.get(page_url, headers=reviewer_headers).text
    assert 'data-action="review"' in reviewer_page
    assert 'data-action="return"' in reviewer_page
    assert 'id="return-comment"' in reviewer_page
    assert client.post(f"/api/evaluations/{evaluation['id']}/workflow/review", headers=reviewer_headers).status_code == 200
    assert 'data-action="return"' not in client.get(page_url, headers=reviewer_headers).text
    approver_headers = {"Authorization": "Bearer approve-token"}
    approver_page = client.get(page_url, headers=approver_headers).text
    assert 'data-action="approve"' in approver_page
    assert 'data-action="return"' in approver_page
    assert 'id="return-comment"' in approver_page
    assert client.post(f"/api/evaluations/{evaluation['id']}/workflow/return", headers=reviewer_headers).status_code == 403


def test_segregation_of_duties_uses_stable_accounts_not_display_names(tmp_path):
    app = create_app(f"sqlite:///{tmp_path / 'stable-actors.db'}", seed=False, auth_tokens={
        "submit-token": {"name": "Shared Display Name", "roles": ["ADMIN", "EVALUATOR"]},
        "review-token": {"name": "Shared Display Name", "roles": ["REVIEWER"]},
    })
    submitter = TestClient(app, headers={"Authorization": "Bearer submit-token"})
    supplier = submitter.post("/api/suppliers", json={"name": "Stable Actor Supplier"}).json()
    contract = submitter.post("/api/contracts", json={"supplier_id": supplier["id"], "contract_number": "ACTOR-1", "procurement_type": "Construction", "region": "Atlantic"}).json()
    evaluation = submitter.post("/api/evaluations", json={"contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator", "scores": {"Quality of Workmanship": 15, "Time": 15, "Project Management": 15, "Contract Management": 15, "Health and Safety": 15}}).json()
    assert submitter.post(f"/api/evaluations/{evaluation['id']}/workflow/submit").status_code == 200
    reviewer = TestClient(app, headers={"Authorization": "Bearer review-token"})
    assert reviewer.post(f"/api/evaluations/{evaluation['id']}/workflow/review").status_code == 200


def test_create_evaluation_navigation_is_role_aware(tmp_path):
    client = make_client(tmp_path)
    assert '/evaluations/new' in client.get("/").text
    reviewer_headers = {"Authorization": "Bearer review-token"}
    reviewer_page = client.get("/", headers=reviewer_headers).text
    assert '/evaluations/new' not in reviewer_page
    assert '/evaluations/new' not in client.get("/evaluations", headers=reviewer_headers).text
    assert client.get("/evaluations/new", headers=reviewer_headers).status_code == 403


def test_new_evaluation_requires_deliberate_score_entry(tmp_path):
    client = make_client(tmp_path)
    page = client.get("/evaluations/new").text
    assert 'data-score="Quality of Workmanship" value="15"' not in page
    assert 'data-score="Quality of Workmanship" required' in page
    assert 'id="score-preview">Incomplete<' in page
    assert 'aria-live="polite"' in page


def test_submitted_evaluation_cannot_be_edited_until_returned(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Workflow Lock Supplier"}).json()
    contract = client.post("/api/contracts", json={"supplier_id": supplier["id"], "contract_number": "LOCK-1", "procurement_type": "Construction", "region": "Atlantic", "status": "ACTIVE"}).json()
    evaluation = client.post("/api/evaluations", json={"contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator", "scores": {"Quality of Workmanship": 15, "Time": 15, "Project Management": 15, "Contract Management": 15, "Health and Safety": 15}}).json()

    assert client.post(f"/api/evaluations/{evaluation['id']}/workflow/submit").status_code == 200
    locked = client.patch(f"/api/evaluations/{evaluation['id']}", json={"comments": "Changed after submission."})
    assert locked.status_code == 409
    assert "returned" in locked.json()["detail"].lower()

    assert client.post(
        f"/api/evaluations/{evaluation['id']}/workflow/return",
        headers={"Authorization": "Bearer review-token"},
        json={"comment": "Please revise the evaluation based on the review findings."},
    ).status_code == 200
    editable = client.patch(f"/api/evaluations/{evaluation['id']}", json={"comments": "Changed after return."})
    assert editable.status_code == 200


def test_return_requires_review_comments_and_enables_evaluator_revision(tmp_path):
    client = make_client(tmp_path)
    reviewer = {"Authorization": "Bearer review-token"}
    approver = {"Authorization": "Bearer approve-token"}
    supplier = client.post("/api/suppliers", json={"name": "Returned Evaluation Supplier"}).json()
    contract = client.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "RETURN-EDIT-1",
        "procurement_type": "Construction", "region": "Atlantic", "status": "ACTIVE"
    }).json()
    evaluation = client.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Original Evaluator",
        "evaluation_date": "2026-07-19", "scores": {
            "Quality of Workmanship": 15, "Time": 15, "Project Management": 15,
            "Contract Management": 15, "Health and Safety": 15,
        }, "comments": "Original evaluation evidence."
    }).json()
    evaluation_id = evaluation["id"]
    assert client.post(f"/api/evaluations/{evaluation_id}/workflow/submit").status_code == 200

    missing_comment = client.post(f"/api/evaluations/{evaluation_id}/workflow/return", headers=reviewer)
    assert missing_comment.status_code == 422
    assert "comment" in missing_comment.json()["detail"].lower()

    reviewer_comment = "Please revise the Time score using the accepted completion date."
    returned = client.post(
        f"/api/evaluations/{evaluation_id}/workflow/return",
        headers=reviewer,
        json={"comment": reviewer_comment},
    )
    assert returned.status_code == 200, returned.text
    assert returned.json()["status"] == "RETURNED"
    history = client.get(f"/api/evaluations/{evaluation_id}/history").json()
    assert history["return_comments"][-1]["comment"] == reviewer_comment
    assert history["return_comments"][-1]["user"].startswith("Independent Reviewer")

    detail = client.get(f"/evaluations/{evaluation_id}/view")
    assert reviewer_comment in detail.text
    assert f'href="/evaluations/{evaluation_id}/edit"' in detail.text
    edit_page = client.get(f"/evaluations/{evaluation_id}/edit")
    assert edit_page.status_code == 200
    assert 'data-evaluation-id="' + str(evaluation_id) + '"' in edit_page.text
    assert 'value="15.0"' in edit_page.text
    assert "Original evaluation evidence." in edit_page.text
    assert reviewer_comment in edit_page.text
    assert client.get(f"/evaluations/{evaluation_id}/edit", headers=reviewer).status_code == 403

    revised = client.patch(f"/api/evaluations/{evaluation_id}", json={
        "scores": {
            "Quality of Workmanship": 15, "Time": 12, "Project Management": 15,
            "Contract Management": 15, "Health and Safety": 15,
        },
        "comments": "Revised using the accepted completion date as requested by the reviewer.",
    })
    assert revised.status_code == 200, revised.text
    assert revised.json()["scores"]["Time"] == 12

    assert client.post(f"/api/evaluations/{evaluation_id}/workflow/submit").status_code == 200
    assert client.post(f"/api/evaluations/{evaluation_id}/workflow/review", headers=reviewer).status_code == 200
    approver_comment = "Please clarify the supporting evidence for the Contract Management score."
    approver_return = client.post(
        f"/api/evaluations/{evaluation_id}/workflow/return",
        headers=approver,
        json={"comment": approver_comment},
    )
    assert approver_return.status_code == 200, approver_return.text
    history = client.get(f"/api/evaluations/{evaluation_id}/history").json()
    assert [item["comment"] for item in history["return_comments"]] == [reviewer_comment, approver_comment]


def test_only_assigned_evaluator_may_revise_a_returned_evaluation(tmp_path):
    app = create_app(f"sqlite:///{tmp_path / 'assigned-evaluator.db'}", seed=False, auth_tokens={
        "admin": {"name": "Administrator", "roles": ["ADMIN"]},
        "evaluator-a": {"name": "Evaluator A", "roles": ["EVALUATOR"]},
        "evaluator-b": {"name": "Evaluator B", "roles": ["EVALUATOR"]},
        "reviewer": {"name": "Independent Reviewer", "roles": ["REVIEWER"]},
    })
    admin = TestClient(app, headers={"Authorization": "Bearer admin"})
    evaluator_a = TestClient(app, headers={"Authorization": "Bearer evaluator-a"})
    supplier = admin.post("/api/suppliers", json={"name": "Assigned Supplier"}).json()
    contract = admin.post("/api/contracts", json={
        "supplier_id": supplier["id"], "contract_number": "ASSIGNED-1",
        "procurement_type": "Construction", "region": "Atlantic",
    }).json()
    evaluation = evaluator_a.post("/api/evaluations", json={
        "contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator A",
        "scores": {name: 15 for name in CONSTRUCTION_CRITERIA},
    }).json()
    evaluation_id = evaluation["id"]
    assert evaluator_a.post(f"/api/evaluations/{evaluation_id}/workflow/submit").status_code == 200
    assert evaluator_a.post(
        f"/api/evaluations/{evaluation_id}/workflow/return",
        headers={"Authorization": "Bearer reviewer"},
        json={"comment": "Please correct the supporting completion-date evidence."},
    ).status_code == 200

    evaluator_b_headers = {"Authorization": "Bearer evaluator-b"}
    assert evaluator_a.get(f"/evaluations/{evaluation_id}/edit", headers=evaluator_b_headers).status_code == 403
    assert evaluator_a.patch(
        f"/api/evaluations/{evaluation_id}", headers=evaluator_b_headers,
        json={"comments": "Unauthorized revision attempt."},
    ).status_code == 403
    assert evaluator_a.post(
        f"/api/evaluations/{evaluation_id}/workflow/submit", headers=evaluator_b_headers,
    ).status_code == 403
    assert evaluator_a.get(f"/evaluations/{evaluation_id}/edit").status_code == 200


def test_xlsx_neutralizes_formula_injection():
    content = performance_history_xlsx([{"supplier": "=1+1", "contract": "+CMD", "evaluator": "@payload"}])
    ws = load_workbook(BytesIO(content), data_only=False).active
    assert ws["A2"].data_type != "f"
    assert ws["A2"].value == "'=1+1"
    assert ws["B2"].data_type != "f"


def test_rbac_regime_dates_sod_and_attachment_controls(tmp_path):
    client = make_client(tmp_path)
    supplier = client.post("/api/suppliers", json={"name": "Controlled Supplier"}).json()
    contract = client.post("/api/contracts", json={"supplier_id": supplier["id"], "contract_number": "CTRL-1", "procurement_type": "Construction", "region": "Atlantic", "status": "ACTIVE"}).json()
    construction = {"contract_id": contract["id"], "model": "CONSTRUCTION", "evaluator": "Evaluator", "evaluation_date": "2026-07-17", "scores": {"Quality of Workmanship": 15, "Time": 15, "Project Management": 15, "Contract Management": 15, "Health and Safety": 15}}
    ae = construction | {"model": "AE", "scores": {"Quality of Design": 80, "Quality of Deliverables": 80, "Technical Competence": 80, "Contract Administration": 80, "Project Management": 80, "Schedule Compliance": 80, "Cost Control": 80, "Client Service": 80}}
    assert client.post("/api/evaluations", json=ae).status_code == 422
    assert client.post("/api/evaluations", json=construction, headers={"Authorization": "Bearer review-token"}).status_code == 403
    evaluation = client.post("/api/evaluations", json=construction).json()
    assert client.patch(f"/api/evaluations/{evaluation['id']}", json={"due_date": "2026-07-01"}).status_code == 422
    spoofed = client.post(f"/api/evaluations/{evaluation['id']}/attachments", files={"file": ("evidence.pdf", b"not a pdf", "application/pdf")})
    assert spoofed.status_code == 415
    assert client.post(f"/api/evaluations/{evaluation['id']}/workflow/submit").status_code == 200
    assert client.post(f"/api/evaluations/{evaluation['id']}/workflow/review").status_code == 409
    locked_attachment = client.post(f"/api/evaluations/{evaluation['id']}/attachments", files={"file": ("evidence.txt", b"new evidence", "text/plain")})
    assert locked_attachment.status_code == 409
    exported = client.get("/api/reports/performance-history.xlsx")
    assert load_workbook(BytesIO(exported.content)).active.max_row == 1
